"""Synthetic-only Word fallback tests: no user files or external model calls."""
from io import BytesIO
import os
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument

from apps.api.models.blobs import Blob
from apps.api.models.documents import Document, DocumentSourceType, DocumentVersion
from apps.api.models.processing import ProcessingJob
from apps.api.parsers.base import Block, ParserResult, StructuredContent
from apps.api.parsers.docx import DocxParser
from apps.api.parsers.pdf import PdfOcrOptions
from apps.worker.services import processor as proc
from apps.worker.tests.test_processor import engine, session  # noqa: F401


def word_result(text="", images=1, kind="docx"):
    return ParserResult(True, StructuredContent(kind,
        [Block("paragraph", text, [])] if text else [],
        metadata={"docx_extraction": {"embedded_image_count": images}}))


def test_docx_counts_embedded_images_not_pages():
    # python-docx deduplicates repeated image parts: never use this as page count.
    from PIL import Image
    image = BytesIO()
    Image.new("RGB", (20, 20), "white").save(image, format="PNG")
    doc = WordDocument()
    doc.add_picture(BytesIO(image.getvalue()))
    doc.add_picture(BytesIO(image.getvalue()))
    raw = BytesIO()
    doc.save(raw)
    result = DocxParser().parse(raw.getvalue())
    assert result.success
    assert not result.structured_content.full_text()
    assert result.structured_content.metadata["docx_extraction"]["embedded_image_count"] == 1


def test_empty_word_thumbnail_is_not_a_body_image():
    buffer = BytesIO()
    WordDocument().save(buffer)
    result = DocxParser().parse(buffer.getvalue())
    assert result.structured_content.metadata["docx_extraction"]["embedded_image_count"] == 0


@pytest.mark.skipif(os.environ.get("CANGZHI_TEST_WORD_OCR_INTEGRATION") != "1",
                    reason="requires isolated LibreOffice/Tesseract runtime")
@pytest.mark.parametrize("kind", ["docx", "wps"])
def test_real_word_preview_and_local_ocr(session, tmp_path, monkeypatch, kind):
    from PIL import Image, ImageDraw, ImageFont
    from docx.shared import Inches

    image = Image.new("RGB", (1200, 300), "white")
    ImageDraw.Draw(image).text((50, 80), "CANGZHI OCR TEST 12345",
                              font=ImageFont.truetype("LiberationSans-Regular.ttf", 48), fill="black")
    picture = BytesIO()
    image.save(picture, format="PNG")
    word = WordDocument()
    word.add_picture(BytesIO(picture.getvalue()), width=Inches(6))
    buffer = BytesIO()
    word.save(buffer)
    source = buffer.getvalue()
    document = Document(title="synthetic", source_type=DocumentSourceType.file)
    session.add(document)
    session.flush()
    version = DocumentVersion(document_id=document.id, version_number=1, content_hash="synthetic")
    session.add(version)
    session.flush()
    monkeypatch.setattr(proc.settings, "storage_path", str(tmp_path))
    if kind == "wps":
        from apps.api.parsers.wps import WpsParser
        native = WpsParser().parse(source)  # OOXML-compatible WPS container
    else:
        native = DocxParser().parse(source)
    assert not native.structured_content.full_text()
    result = proc._word_ocr_fallback(session, version, native, source,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"synthetic.{kind}",
        PdfOcrOptions(language="eng", external_max_pages=0))
    assert result.success, result.error_message
    assert "CANGZHI OCR TEST 12345" in result.structured_content.full_text()
    assert version.preview_blob_id is not None
    extraction = result.structured_content.metadata["pdf_extraction"]
    assert extraction["page_count"] == 1
    assert extraction["ocr_completed_pages"] == [1]
    assert not extraction.get("external_attempted_pages")
    assert result.structured_content.document_type == kind
    if kind == "wps":
        assert result.structured_content.metadata["source_format"] == "wps"
    assert buffer.getvalue() == source


@pytest.mark.skipif(os.environ.get("CANGZHI_TEST_WORD_OCR_INTEGRATION") != "1",
                    reason="requires isolated LibreOffice runtime")
def test_real_word97_compatible_wps_parse_and_preview(session, tmp_path, monkeypatch):
    import subprocess
    from apps.api.parsers.wps import WpsParser
    word = WordDocument()
    word.add_heading("Compatibility test", 1)
    word.add_paragraph("CANGZHI WPS FACT 12345")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "TABLE FACT 67890"
    table.cell(1, 1).text = "123"
    # Keep an explicit trailing paragraph: Word97 fixture export itself can
    # flatten a terminal table. This is not a claim of lossless Office export.
    word.add_paragraph("End of compatibility fixture")
    path = tmp_path / "synthetic.docx"
    word.save(path)
    subprocess.run(["soffice", f"-env:UserInstallation={(tmp_path / 'export-profile').as_uri()}",
                    "--headless", "--convert-to", "wps:MS Word 97", "--outdir", str(tmp_path), str(path)],
                   check=True, capture_output=True, timeout=60)
    source = path.with_suffix(".wps").read_bytes()
    assert source.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    parsed = WpsParser().parse(source)
    assert parsed.success, parsed.error_message
    assert "CANGZHI WPS FACT 12345" in parsed.structured_content.full_text()
    assert "TABLE FACT 67890" in parsed.structured_content.full_text()
    assert any(block.type == "table" for block in parsed.structured_content.blocks), [(block.type, block.text) for block in parsed.structured_content.blocks]
    assert parsed.structured_content.document_type == "wps"
    doc = Document(title="synthetic WPS", source_type=DocumentSourceType.file)
    session.add(doc)
    session.flush()
    version = DocumentVersion(document_id=doc.id, version_number=1, content_hash="wps-synthetic")
    session.add(version)
    session.flush()
    monkeypatch.setattr(proc.settings, "storage_path", str(tmp_path / "storage"))
    assert proc._ensure_word_pdf_preview(session, version, source, "application/kswps", "synthetic.WPS")
    blob = session.get(Blob, version.preview_blob_id)
    assert proc._read_blob_bytes(blob).startswith(b"%PDF-")
    assert path.with_suffix(".wps").read_bytes() == source


@pytest.mark.parametrize("reason", ["wps_unsupported_container", "wps_converter_unavailable",
    "wps_conversion_failed", "wps_conversion_timeout", "wps_conversion_error"])
def test_wps_failures_do_not_automatically_repeat_conversion(reason):
    assert proc._is_terminal_parse_failure({"reason": reason})


@pytest.mark.parametrize("original", [word_result("正文"), word_result(kind="pdf"), ParserResult(False)])
def test_non_candidates_do_not_convert_or_ocr(monkeypatch, original):
    def unexpected(*args, **kwargs):
        pytest.fail("must not convert or OCR")
    monkeypatch.setattr(proc, "_ensure_word_pdf_preview", unexpected)
    monkeypatch.setattr(proc, "get_parser_for_content", unexpected)
    assert proc._word_ocr_fallback(None, None, original, b"source", "", "x.docx", PdfOcrOptions()) is original


@pytest.mark.parametrize("images,enabled,reason", [
    (0, True, "word_empty_text"), (1, False, "word_ocr_disabled"),
])
def test_empty_and_disabled_fail_without_external_calls(monkeypatch, images, enabled, reason):
    monkeypatch.setattr(proc, "_ensure_word_pdf_preview", lambda *args: pytest.fail("no conversion"))
    result = proc._word_ocr_fallback(None, None, word_result(images=images), b"", "", "x.docx", PdfOcrOptions(enabled=enabled))
    assert not result.success
    assert result.error_details["reason"] == reason
    assert proc._is_terminal_parse_failure(result.error_details)


@pytest.mark.parametrize("kind", ["doc", "docx", "wps"])
@pytest.mark.parametrize("outcome", ["partial", "empty", "failed", "exception", "preview_failed"])
def test_preview_reuse_and_failure_boundaries(monkeypatch, kind, outcome):
    source = b"original Word unchanged"
    preview = SimpleNamespace(id=42)
    version = SimpleNamespace(preview_blob_id=42)
    session = SimpleNamespace(get=lambda model, id: preview)
    options = PdfOcrOptions(external_max_pages=2, max_pages=3)
    calls = []
    extraction = {"ocr_status": "partial", "page_count": 4,
                  "quality_policy": "page-quarantine-v1", "low_quality_pages": [2],
                  "quarantined_pages": [2], "ocr_completed_pages": [1], "ocr_skipped_pages": [3, 4]}

    def ensure(*args):
        assert args[2] == source
        return outcome != "preview_failed"

    def parse(content, mime):
        assert content == b"exact preview bytes"
        assert mime == "application/pdf"
        calls.append("parse")
        if outcome == "exception":
            raise RuntimeError("upstream secret must not be persisted")
        if outcome == "failed":
            return ParserResult(False, error_message="private upstream response")
        return ParserResult(True, StructuredContent("pdf",
            [] if outcome == "empty" else [Block("paragraph", "已识别正文", [], page=1)],
            metadata={"pdf_extraction": extraction}))

    def factory(mime, filename, *, pdf_ocr_options):
        assert pdf_ocr_options is options  # same local/external limits and provider
        return SimpleNamespace(parse=parse)

    monkeypatch.setattr(proc, "_ensure_word_pdf_preview", ensure)
    monkeypatch.setattr(proc, "_read_blob_bytes", lambda blob: b"exact preview bytes")
    monkeypatch.setattr(proc, "get_parser_for_content", factory)
    result = proc._word_ocr_fallback(session, version, word_result(kind=kind), source, "", f"x.{kind}", options)
    if outcome == "partial":
        assert result.success
        structured = result.structured_content
        assert structured.document_type == kind
        assert structured.metadata["pdf_extraction"] == extraction
        assert structured.metadata["word_ocr"]["preview_blob_id"] == 42
        assert structured.blocks[0].page == 1
        assert structured.blocks[0].paragraph_index is None
        assert "docx_body_index" not in structured.blocks[0].extra
    else:
        assert not result.success
        assert proc._is_terminal_parse_failure(result.error_details)
        assert "secret" not in str(result)
        assert "private upstream" not in str(result)
    assert len(calls) == (0 if outcome == "preview_failed" else 1)


@pytest.mark.parametrize("recognized", [False, True])
@pytest.mark.parametrize("kind", ["doc", "wps"])
def test_pipeline_never_completes_empty_word(session, monkeypatch, recognized, kind):
    doc = Document(title="synthetic image Word", source_type=DocumentSourceType.file)
    session.add(doc)
    session.flush()
    blob = Blob(sha256="source", storage_key="synthetic", content_type="application/msword",
                original_filename=f"test.{kind}", file_size=10)
    session.add(blob)
    session.flush()
    version = DocumentVersion(document_id=doc.id, version_number=1, content_hash="source",
                              blob_id=blob.id, processing_status="created",
                              raw_content="previous extraction")
    session.add(version)
    session.flush()
    doc.current_version_id = version.id
    job = ProcessingJob(document_id=doc.id, document_version_id=version.id,
                        stage="parsing", config_version="1", idempotency_key="word-ocr-test")
    session.add(job)
    session.commit()
    monkeypatch.setattr(proc, "_load_content_for_parsing", lambda *args: (b"source", "application/msword", f"test.{kind}"))
    monkeypatch.setattr(proc, "_build_pdf_ocr_options", lambda *args: PdfOcrOptions())
    def ensure(*args):
        version.preview_blob_id = blob.id  # synthetic stand-in, never a real source
        return True
    monkeypatch.setattr(proc, "_ensure_word_pdf_preview", ensure)
    monkeypatch.setattr(proc, "_read_blob_bytes", lambda *args: b"preview")
    def factory(mime, *args, **kwargs):
        result = word_result(kind=kind) if mime != "application/pdf" else ParserResult(
            True, StructuredContent("pdf", [Block("paragraph", "recognized text", [], page=1)] if recognized else [],
                                    metadata={"pdf_extraction": {"ocr_status": "completed" if recognized else "failed"}}))
        return SimpleNamespace(parse=lambda *args: result)
    monkeypatch.setattr(proc, "get_parser_for_content", factory)
    assert proc.process_single_job(session, job.id) is recognized
    session.refresh(job)
    session.refresh(version)
    assert version.blob_id == blob.id
    assert version.content_hash == "source"
    downstream = session.query(ProcessingJob).filter(ProcessingJob.id != job.id).all()
    if recognized:
        assert job.status == "completed"
        assert version.raw_content == "recognized text"
        assert {j.stage for j in downstream} >= {"chunking", "understanding"}
    else:
        assert job.status == "failed"
        assert version.processing_status == "failed"
        assert version.raw_content == "previous extraction"
        assert job.next_retry_at is None
        assert not downstream
        assert job.error_details["reason"] == "word_ocr_empty"

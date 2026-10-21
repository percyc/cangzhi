"""Compatible WPS containers, not a claim about every Kingsoft native version."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import subprocess

import pytest
from docx import Document

from apps.api.api.files import validate_file_type
from apps.api.api.webdav import DEFAULT_EXTENSIONS
from apps.api.parsers import WpsParser, get_parser_for_content

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@pytest.mark.parametrize("mime", ["application/octet-stream", "application/kswps", "application/msword",
                                  "application/vnd.ms-works", "application/wps", "application/x-wps",
                                  "application/vnd.wps-office.wps"])
def test_wps_upload_preserves_original_and_download(client, mime):
    api, _ = client
    source = OLE + b"synthetic compatibility fixture"
    result = api.post("/api/files/upload", files={"file": ("材料.WPS", source, mime)})
    assert result.status_code == 201, result.text
    body = result.json()
    assert body["current_version"]["blob"]["original_filename"] == "材料.WPS"
    assert api.get(f"/api/documents/{body['id']}/original").content == source
    assert isinstance(get_parser_for_content(mime, "材料.WPS"), WpsParser)


def test_wps_validation_and_dispatch():
    from fastapi import HTTPException
    from apps.api.parsers.doc import DocParser
    assert ".wps" in DEFAULT_EXTENSIONS
    assert isinstance(get_parser_for_content("application/kswps"), WpsParser)
    assert type(get_parser_for_content("application/msword", "x.doc")) is DocParser
    with pytest.raises(HTTPException):
        validate_file_type("x.wps", "application/pdf")
    for extension in ("et", "dps"):
        with pytest.raises(HTTPException):
            validate_file_type(f"x.{extension}", None)


def test_wps_evidence_uses_word_preview_path():
    from apps.api.services.evidence import _classify, _preview_url
    from apps.api.models.documents import DocumentSourceType
    kind = _classify(source_type=DocumentSourceType.file, document_type="wps", table_location={})
    assert kind == "pdf_word"
    assert _preview_url(1, 2, kind) == "/api/documents/1/preview?version_id=2"


def test_wps_conversion_preserves_structure_and_source(monkeypatch):
    source = OLE + b"synthetic"
    calls = []
    def convert(command, **options):
        path = Path(command[-1])
        assert path.name == "source.wps"
        assert path.read_bytes() == source
        assert options["timeout"] == 60
        assert "--infilter=MS Word 97" in command
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        doc = Document()
        doc.add_heading("兼容标题", 1)
        doc.add_paragraph("不可改写的正文")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格事实"
        doc.save(path.with_suffix(".docx"))
        calls.append(path)
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr("apps.api.parsers.doc.shutil.which", lambda _: "/usr/bin/soffice")
    monkeypatch.setattr("apps.api.parsers.doc.subprocess.run", convert)
    result = WpsParser().parse(source)
    assert result.success
    structured = result.structured_content
    assert structured.document_type == "wps"
    assert [b.type for b in structured.blocks] == ["heading", "paragraph", "table"]
    assert "不可改写的正文" in structured.full_text()
    assert structured.metadata["source_format"] == "wps"
    assert structured.metadata["converted_format"] == "docx"
    assert not calls[0].exists()  # temporary conversion and profile removed


@pytest.mark.parametrize("failure,reason", [
    ("missing", "wps_converter_unavailable"), ("failed", "wps_conversion_failed"),
    ("timeout", "wps_conversion_timeout"), ("oserror", "wps_conversion_error"),
])
def test_wps_failures_are_actionable_without_leaking_converter_output(monkeypatch, failure, reason):
    monkeypatch.setattr("apps.api.parsers.doc.shutil.which", lambda _: None if failure == "missing" else "/soffice")
    def convert(*args, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired("soffice", 60)
        if failure == "oserror":
            raise OSError("private source text")
        return SimpleNamespace(returncode=1, stderr=b"private source text", stdout=b"")
    monkeypatch.setattr("apps.api.parsers.doc.subprocess.run", convert)
    result = WpsParser().parse(OLE + b"source")
    assert not result.success
    assert result.error_details == {"reason": reason}
    assert "private" not in str(result)
    assert "WPS" in result.error_message


@pytest.mark.parametrize("source", [b"", b"random binary\x00\xff", b"plain text masquerading as WPS"])
def test_unknown_wps_never_falls_back_to_text(monkeypatch, source):
    monkeypatch.setattr("apps.api.parsers.doc.subprocess.run", lambda *a, **k: pytest.fail("no conversion"))
    result = WpsParser().parse(source)
    assert not result.success
    assert result.error_details["reason"] == "wps_unsupported_container"

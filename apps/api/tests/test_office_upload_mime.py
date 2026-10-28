"""Platform-independent upload headers, preserved originals and parser dispatch."""
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import HTTPException

from apps.api.api.files import validate_file_type
from apps.api.parsers import DocxParser, get_parser_for_content
from apps.api.services.webdav import RemoteEntry

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MIMES = ["application/zip", "application/x-zip", "application/x-zip-compressed"]


def docx_bytes():
    document = Document()
    document.add_heading("合成兼容性样例", level=1)
    document.add_paragraph("原始正文不可改写。")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格事实"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("mime", [DOCX_MIME, *ZIP_MIMES, "application/octet-stream",
                                  "application/wps-office.docx",
                                  " Application/WPS-Office.Docx ; charset=binary",
                                  None, "", " Application/ZIP ; charset=binary",
                                  DOCX_MIME.upper() + "; charset=utf-8"])
def test_docx_upload_download_and_parse(client, mime):
    api, _ = client
    source = docx_bytes()
    response = api.post("/api/files/upload", files={"file": ("材料.DOCX", source, mime)})
    assert response.status_code == 201, response.text
    body = response.json()
    blob = body["current_version"]["blob"]
    if mime and "wps-office.docx" in mime.lower():
        assert blob["content_type"] == DOCX_MIME
    downloaded = api.get(f"/api/documents/{body['id']}/original")
    assert downloaded.status_code == 200
    assert downloaded.content == source
    parser = get_parser_for_content(blob["content_type"], blob["original_filename"])
    assert isinstance(parser, DocxParser)
    result = parser.parse(downloaded.content, blob["content_type"])
    assert result.success
    assert [block.type for block in result.structured_content.blocks] == ["heading", "paragraph", "table"]
    assert result.structured_content.full_text() == "合成兼容性样例\n原始正文不可改写。\n表格事实"


@pytest.mark.parametrize("extension,expected", [
    ("docx", DOCX_MIME),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ("wps", "application/kswps"),
])
@pytest.mark.parametrize("mime", ZIP_MIMES)
def test_container_aliases_normalized(extension, expected, mime):
    assert validate_file_type(f"sample.{extension}", mime) == expected


@pytest.mark.parametrize("mime", [None, "", "   ", " Application/Octet-Stream ; charset=binary"])
def test_generic_mime_compatibility(mime):
    assert validate_file_type("sample.docx", mime) == "application/octet-stream"


@pytest.mark.parametrize("filename,mime", [
    ("sample.zip", "application/zip"), ("sample.exe", "application/octet-stream"),
    ("sample.pdf", "application/zip"), ("sample.doc", "application/zip"),
    ("sample.txt", "application/zip"), ("sample.xls", "application/zip"),
    ("sample.docx", "application/pdf"), ("sample.docx", "text/plain"),
    ("sample.docx", "application/zip-invalid"),
    ("sample.xlsx", "application/wps-office.docx"),
    ("sample.doc", "application/wps-office.docx"),
    ("sample.pdf", "application/wps-office.docx"),
    ("sample.docx", "application/wps-office.unknown"),
])
def test_mismatches_still_rejected_without_storage(client, filename, mime):
    api, storage = client
    response = api.post("/api/files/upload", files={"file": (filename, b"synthetic", mime)})
    assert response.status_code == 415
    assert list(storage._blob_path.iterdir()) == []
    with pytest.raises(HTTPException):
        validate_file_type(filename, mime)


def test_renamed_zip_is_preserved_but_never_parsed_as_text(client):
    api, _ = client
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "Not a Word document")
    source = buffer.getvalue()
    response = api.post("/api/files/upload", files={"file": ("wrong.docx", source, "application/zip")})
    assert response.status_code == 201
    body = response.json()
    blob = body["current_version"]["blob"]
    parser = get_parser_for_content(blob["content_type"], blob["original_filename"])
    assert isinstance(parser, DocxParser)
    assert not parser.parse(source).success
    assert not parser.parse(source[:30]).success
    assert api.get(f"/api/documents/{body['id']}/original").content == source


def test_zip_alias_does_not_bypass_empty_or_size_limits(client, monkeypatch):
    from apps.api.core.config import settings
    api, storage = client
    monkeypatch.setattr(settings, "max_upload_size_mb", 1)
    for source, expected in [(b"", 400), (b"x" * (1024 * 1024 + 1), 413)]:
        response = api.post("/api/files/upload", files={"file": ("sample.docx", source, "application/zip")})
        assert response.status_code == expected
        assert list(storage._blob_path.iterdir()) == []


@pytest.mark.parametrize("mime", [" Application/X-Zip-Compressed ; charset=binary",
                                  "application/wps-office.docx"])
def test_webdav_docx_container_mime(client, monkeypatch, mime):
    api, _ = client
    source_bytes = docx_bytes()
    source = api.post("/api/webdav", json={
        "name": "Synthetic Office source", "base_url": "https://dav.example.com/",
        "username": "reader", "password": "synthetic", "root_path": "/knowledge",
    }).json()

    async def propfind(**kwargs):
        return [RemoteEntry(path="/knowledge/sample.DOCX", is_collection=False, etag='"v1"',
                            last_modified="today", size=len(source_bytes), content_type="application/zip")]

    async def download(**kwargs):
        return source_bytes, mime

    monkeypatch.setattr("apps.api.api.webdav.propfind", propfind)
    monkeypatch.setattr("apps.api.api.webdav.download_file", download)
    monkeypatch.setattr("apps.api.api.documents.download_file", download)
    assert api.post(f"/api/webdav/{source['id']}/scan").status_code == 200
    response = api.post(f"/api/webdav/{source['id']}/sync")
    assert response.status_code == 200
    assert response.json()["imported"] == 1, response.text
    entry = api.get(f"/api/webdav/{source['id']}/entries").json()[0]
    assert api.get(f"/api/documents/{entry['document_id']}/original").content == source_bytes

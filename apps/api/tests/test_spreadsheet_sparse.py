from io import BytesIO
import re
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from apps.api.parsers import spreadsheet as parser


def workbook_bytes(edit):
    book = Workbook()
    edit(book.active)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def rewrite_sheet(content, transform):
    stream = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(stream, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = transform(data.decode()).encode()
            target.writestr(item, data)
    return stream.getvalue()


def test_phantom_format_extent_no_longer_consumes_cell_budget(monkeypatch):
    def edit(sheet):
        sheet.append(["名称", "数量"])
        sheet.append(["甲", 3])
        sheet["GF21460"].fill = PatternFill("solid", fgColor="FFFFFF")
    monkeypatch.setattr(parser, "MAX_CELLS_TOTAL", 10)
    result = parser.XlsxParser().parse(workbook_bytes(edit))
    assert result.success, result.error_message
    meta = result.structured_content.metadata
    assert meta["non_empty_cells"] == 4
    assert meta["sheet_dimensions"][0]["actual_dimension"] == "A1:B2"
    assert meta["sheet_dimensions"][0]["declared_dimension"] == "A1:GF21460"


def test_underreported_dimensions_and_absolute_coordinates():
    def edit(sheet):
        sheet["C5"] = "首个证据"
        sheet["E7"] = "尾部证据"
    content = rewrite_sheet(workbook_bytes(edit), lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1"', xml))
    result = parser.XlsxParser().parse(content)
    assert result.success, result.error_message
    text = result.structured_content.full_text()
    assert "行 5｜C列=首个证据" in text
    assert "行 7｜E列=尾部证据" in text
    assert result.structured_content.metadata["data_region_count"] == 2


def test_styled_only_sheet_has_no_content():
    def edit(sheet):
        sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FFFFFF")
    result = parser.XlsxParser().parse(workbook_bytes(edit))
    assert result.success
    assert result.structured_content.blocks == []


@pytest.mark.parametrize("coordinate", ["K2", "A11"])
def test_genuine_materialization_and_row_limits_are_not_removed(monkeypatch, coordinate):
    monkeypatch.setattr(parser, "MAX_CELLS_TOTAL", 20)
    monkeypatch.setattr(parser, "MAX_ROWS_PER_SHEET", 10)
    def edit(sheet):
        sheet["A1"] = "开始"
        sheet[coordinate] = "真实值"
    result = parser.XlsxParser().parse(workbook_bytes(edit))
    assert not result.success
    assert result.error_details["reason"] == "spreadsheet_limit_exceeded"


def test_workbook_cumulative_budget_remains_enforced(monkeypatch):
    monkeypatch.setattr(parser, "MAX_CELLS_TOTAL", 6)
    book = Workbook()
    book.active.append([1, 2, 3, 4])
    book.create_sheet().append([1, 2, 3, 4])
    stream = BytesIO()
    book.save(stream)
    result = parser.XlsxParser().parse(stream.getvalue())
    assert not result.success
    assert result.error_details["reason"] == "spreadsheet_limit_exceeded"


def test_formula_cached_value_and_merge_metadata_survive():
    def edit(sheet):
        sheet["A1"] = "标题"
        sheet.merge_cells("A1:C1")
        sheet["C3"] = "=1+2"
    content = rewrite_sheet(workbook_bytes(edit), lambda xml: xml.replace("<f>1+2</f><v></v>", "<f>1+2</f><v>3</v>"))
    result = parser.XlsxParser().parse(content)
    assert result.success, result.error_message
    assert "行 3｜C列=公式=1+2 ⇒ 3" in result.structured_content.full_text()
    assert result.structured_content.metadata["sheet_dimensions"][0]["merge_ranges"] == ["A1:C1"]


def test_invalid_xml_fails_instead_of_silently_returning_empty_success():
    content = workbook_bytes(lambda sheet: sheet.append(["内容"]))
    content = rewrite_sheet(content, lambda xml: xml.replace("</worksheet>", ""))
    result = parser.XlsxParser().parse(content)
    assert not result.success


def test_text_budget_is_independent_of_cell_count(monkeypatch):
    monkeypatch.setattr(parser, "MAX_TEXT_CHARS_TOTAL", 5)
    result = parser.XlsxParser().parse(workbook_bytes(lambda sheet: sheet.append(["abcdef"])))
    assert not result.success
    assert result.error_details["reason"] == "spreadsheet_limit_exceeded"


def test_row_iterators_close_on_cancel():
    from types import SimpleNamespace
    closed = []
    class Sheet:
        def iter_rows(self, **kwargs):
            try:
                yield [SimpleNamespace(value="text", number_format=None)]
                yield [SimpleNamespace(value="more", number_format=None)]
            finally:
                closed.append(True)
    rows = parser._xlsx_rows(Sheet(), Sheet(), min_row=1, max_row=2, min_col=1, max_col=1)
    next(rows)
    rows.close()
    assert len(closed) == 2

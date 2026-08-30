from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from apps.api.parsers import (
    DocParser,
    DocxParser,
    MarkdownParser,
    NoteParser,
    PdfOcrOptions,
    PdfParser,
    TextParser,
    XlsParser,
    XlsxParser,
    get_parser_for_content,
)
from apps.api.parsers import spreadsheet as spreadsheet_parser


class TestNoteParser:
    def test_parse_empty_note(self):
        parser = NoteParser()
        result = parser.parse("")
        assert result.success is True
        assert result.structured_content is not None
        assert len(result.structured_content.blocks) == 0

    def test_parse_simple_note(self):
        parser = NoteParser()
        text = """# Title
## Section
This is a paragraph.
Another paragraph.
"""
        result = parser.parse(text)
        assert result.success is True
        blocks = result.structured_content.blocks
        assert len(blocks) == 4
        assert blocks[0].type == "heading"
        assert blocks[0].level == 1
        assert blocks[0].text == "Title"
        assert blocks[1].type == "heading"
        assert blocks[1].level == 2
        assert blocks[1].text == "Section"

    def test_parse_note_full_text(self):
        parser = NoteParser()
        text = "Hello\nWorld"
        result = parser.parse(text)
        assert result.success is True
        full_text = result.structured_content.full_text()
        assert full_text == "Hello\nWorld"


class TestTextParser:
    def test_decode_utf8(self):
        parser = TextParser()
        content = "Hello 中文".encode("utf-8")
        result = parser.parse(content)
        assert result.success is True
        assert len(result.structured_content.blocks) == 1

    def test_decode_utf8_bom(self):
        parser = TextParser()
        content = b'\xef\xbb\xbfHello'
        result = parser.parse(content)
        assert result.success is True
        assert result.structured_content.blocks[0].text == "Hello"

    def test_decode_error_handling(self):
        parser = TextParser()
        content = b'\xff\xfe\xfd'
        result = parser.parse(content)
        assert result.success is False
        assert "无法解码" in result.error_message


class TestMarkdownParser:
    def test_parse_headings(self):
        parser = MarkdownParser()
        content = """# Heading 1
## Heading 2
### Heading 3
""".encode("utf-8")
        result = parser.parse(content)
        assert result.success is True
        blocks = result.structured_content.blocks
        assert len(blocks) == 3
        assert blocks[0].level == 1
        assert blocks[1].level == 2
        assert blocks[2].level == 3

    def test_parse_code_block(self):
        parser = MarkdownParser()
        content = """```python
def hello():
    print("hello")
```
""".encode("utf-8")
        result = parser.parse(content)
        assert result.success is True
        blocks = result.structured_content.blocks
        assert len(blocks) == 1
        assert blocks[0].type == "code_block"
        assert "def hello" in blocks[0].text

    def test_parse_list_items(self):
        parser = MarkdownParser()
        content = """- Item 1
- Item 2
1. Item 3
""".encode("utf-8")
        result = parser.parse(content)
        assert result.success is True
        blocks = result.structured_content.blocks
        assert len(blocks) == 3
        assert all(b.type == "list_item" for b in blocks)

    def test_get_parser_for_markdown_by_extension(self):
        parser = get_parser_for_content("text/plain", "test.md")
        assert isinstance(parser, MarkdownParser)


class TestStructuredContent:
    def test_to_dict_has_correct_schema(self):
        from apps.api.parsers.base import Block, StructuredContent
        blocks = [Block(type="paragraph", text="test", heading_path=[])]
        sc = StructuredContent(document_type="txt", blocks=blocks)
        data = sc.to_dict()
        assert data["schema_version"] == 1
        assert "document_type" in data
        assert len(data["blocks"]) == 1
        assert "heading_path" in data["blocks"][0]


class TestOfficeParsers:
    def test_spreadsheet_limits_accept_current_large_personal_tables(self):
        assert spreadsheet_parser.MAX_ROWS_PER_SHEET >= 102_966
        assert spreadsheet_parser.MAX_CELLS_TOTAL >= 1_750_422

    def test_spreadsheet_row_limit_remains_enforced(self, monkeypatch):
        monkeypatch.setattr(spreadsheet_parser, "MAX_ROWS_PER_SHEET", 3)

        with pytest.raises(
            spreadsheet_parser.SpreadsheetLimitError,
            match="超过 3 行限制",
        ):
            spreadsheet_parser._sheet_blocks(
                "大表",
                [["列"]] * 4,
                paragraph_offset=0,
            )

    def test_selects_excel_parsers_by_extension_and_mime(self):
        assert isinstance(
            get_parser_for_content("application/octet-stream", "ledger.xlsx"),
            XlsxParser,
        )
        assert isinstance(
            get_parser_for_content(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                None,
            ),
            XlsxParser,
        )
        assert isinstance(
            get_parser_for_content("application/vnd.ms-excel", "ledger.xls"),
            XlsParser,
        )

    def test_parse_xlsx_preserves_sheets_rows_empty_cells_and_formula(self):
        from openpyxl import Workbook

        workbook = Workbook()
        summary = workbook.active
        summary.title = "汇总"
        summary.append(["项目", "一月", "二月", "合计"])
        summary.append(["收入", 100, None, "=SUM(B2:C2)"])
        summary.append([None, None, None, None])
        details = workbook.create_sheet("明细")
        details.append(["日期", "说明"])
        details.append([datetime(2026, 7, 31, 8, 30), "会员收入"])
        source = BytesIO()
        workbook.save(source)

        result = XlsxParser().parse(source.getvalue())

        assert result.success is True
        structured = result.structured_content
        assert structured.document_type == "xlsx"
        assert structured.metadata["source_format"] == "xlsx"
        assert structured.metadata["sheet_names"] == ["汇总", "明细"]
        assert structured.metadata["non_empty_rows"] == 4
        assert structured.metadata["formula_cells"] == 1
        assert [block.type for block in structured.blocks] == [
            "heading",
            "table",
            "heading",
            "table",
        ]
        assert structured.blocks[1].heading_path == ["汇总", "数据区域 1"]
        assert "行 2｜项目=收入｜一月=100｜合计=公式=SUM(B2:C2)" in structured.blocks[1].text
        assert "行 2｜日期=2026-07-31 08:30:00｜说明=会员收入" in structured.blocks[3].text
        assert structured.blocks[1].extra == {
            "sheet_name": "汇总",
            "region_index": 1,
            "row_start": 2,
            "row_end": 2,
            "header_row": 1,
            "column_names": ["项目", "一月", "二月", "合计"],
        }

    def test_parse_xlsx_splits_regions_and_preserves_display_formats(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "台账"
        sheet.append(["编号", "完成率", "金额"])
        sheet.append([7, 0.25, 1234.5])
        sheet["A2"].number_format = "00000"
        sheet["B2"].number_format = "0.0%"
        sheet["C2"].number_format = "￥#,##0.00"
        sheet.append([None, None, None])
        sheet.append(["张三", "广州"])
        sheet.append(["李四", "深圳"])
        source = BytesIO()
        workbook.save(source)

        result = XlsxParser().parse(source.getvalue())

        assert result.success is True
        structured = result.structured_content
        assert structured.metadata["data_region_count"] == 2
        tables = [block for block in structured.blocks if block.type == "table"]
        assert tables[0].text == (
            "行 2｜编号=00007（原始值：7）｜完成率=25.0%（原始值：0.25）｜"
            "金额=￥1,234.50（原始值：1234.5）"
        )
        assert tables[1].text.startswith("行 4｜A列=张三｜B列=广州")
        assert tables[1].extra["header_row"] is None

    def test_parse_blank_xlsx_returns_no_fake_content(self):
        from openpyxl import Workbook

        workbook = Workbook()
        source = BytesIO()
        workbook.save(source)

        result = XlsxParser().parse(source.getvalue())

        assert result.success is True
        assert result.structured_content.blocks == []
        assert result.structured_content.metadata["sheet_names"] == ["Sheet"]

    def test_selects_legacy_doc_parser_by_extension_and_mime(self):
        assert isinstance(
            get_parser_for_content("application/octet-stream", "legacy.doc"),
            DocParser,
        )
        assert isinstance(
            get_parser_for_content("application/msword", None),
            DocParser,
        )

    def test_doc_parser_reports_missing_converter(self, monkeypatch):
        monkeypatch.setattr(
            "apps.api.parsers.doc.shutil.which",
            lambda _name: None,
        )

        result = DocParser().parse(b"legacy word data")

        assert result.success is False
        assert result.error_details == {"reason": "doc_converter_unavailable"}

    def test_doc_parser_converts_then_preserves_structure(self, monkeypatch):
        from types import SimpleNamespace

        from docx import Document

        def fake_run(arguments, **_kwargs):
            output_directory = arguments[arguments.index("--outdir") + 1]
            converted = Document()
            converted.add_heading("旧版合同", level=1)
            converted.add_paragraph("转换后的正文")
            converted.save(f"{output_directory}/source.docx")
            return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")

        monkeypatch.setattr(
            "apps.api.parsers.doc.shutil.which",
            lambda _name: "/usr/bin/soffice",
        )
        monkeypatch.setattr(
            "apps.api.parsers.doc.subprocess.run",
            fake_run,
        )

        result = DocParser().parse(b"legacy word data")

        assert result.success is True
        assert result.structured_content.document_type == "doc"
        assert result.structured_content.full_text() == "旧版合同\n转换后的正文"
        assert result.structured_content.metadata == {
            "source_format": "doc",
            "converted_format": "docx",
        }

    def test_parse_docx_heading_paragraph_and_table(self):
        from docx import Document

        document = Document()
        document.add_heading("合同条款", level=1)
        document.add_paragraph("这是正文。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲方"
        table.cell(0, 1).text = "乙方"
        source = BytesIO()
        document.save(source)

        result = DocxParser().parse(source.getvalue())
        assert result.success is True
        assert [block.type for block in result.structured_content.blocks] == [
            "heading",
            "paragraph",
            "table",
        ]
        assert result.structured_content.blocks[1].heading_path == ["合同条款"]
        assert result.structured_content.blocks[2].text == "甲方\t乙方"

    def test_parse_docx_with_vendor_specific_style_type(self):
        from docx import Document

        document = Document()
        document.add_heading("非标准样式标题", level=1)
        document.add_paragraph("正文仍应成功提取。")
        source = BytesIO()
        document.save(source)

        malformed = BytesIO()
        with (
            ZipFile(BytesIO(source.getvalue())) as input_archive,
            ZipFile(malformed, "w", ZIP_DEFLATED) as output_archive,
        ):
            for item in input_archive.infolist():
                data = input_archive.read(item.filename)
                if item.filename == "word/styles.xml":
                    data = data.replace(
                        b'w:type="paragraph" w:styleId="Heading1"',
                        b'w:type="titleLevel1" w:styleId="Heading1"',
                    )
                output_archive.writestr(item, data)

        result = DocxParser().parse(malformed.getvalue())

        assert result.success is True
        assert result.structured_content.full_text() == (
            "非标准样式标题\n正文仍应成功提取。"
        )
        assert result.structured_content.blocks[0].type == "heading"

    def test_parse_pdf_preserves_page_count(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        source = BytesIO()
        writer.write(source)

        result = PdfParser().parse(source.getvalue())
        assert result.success is True
        assert result.structured_content.metadata["page_count"] == 1


def _make_pdf_with_image(
    page_size: tuple[float, float] = (200.0, 200.0),
    text_prefix: bytes | None = None,
) -> bytes:
    """Construct a minimal one-page PDF containing a 1x1 raster image.

    The page is otherwise blank, so ``pypdf`` reports zero native text and the
    parser's image probe returns one XObject. ``text_prefix`` lets a single
    test also exercise the "text plus image" branch without writing a real
    font/text object.
    """

    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
        NumberObject,
    )

    width, height = page_size
    writer = PdfWriter()
    page = writer.add_blank_page(width=width, height=height)

    image_stream = DecodedStreamObject()
    image_stream[NameObject("/Type")] = NameObject("/XObject")
    image_stream[NameObject("/Subtype")] = NameObject("/Image")
    image_stream[NameObject("/Width")] = NumberObject(1)
    image_stream[NameObject("/Height")] = NumberObject(1)
    image_stream[NameObject("/ColorSpace")] = NameObject("/DeviceRGB")
    image_stream[NameObject("/BitsPerComponent")] = NumberObject(8)
    image_stream.set_data(b"\xff\x00\x00")
    writer._add_object(image_stream)
    image_ref = image_stream.indirect_reference

    resources = DictionaryObject()
    xobjects = DictionaryObject()
    xobjects[NameObject("/Im0")] = image_ref
    resources[NameObject("/XObject")] = xobjects
    page[NameObject("/Resources")] = resources

    draw_cmd = f"q {width} 0 0 {height} 0 0 cm /Im0 Do Q".encode("latin-1")
    payload = draw_cmd + (text_prefix or b"")
    content_stream = DecodedStreamObject()
    content_stream.set_data(payload)
    writer._add_object(content_stream)
    page[NameObject("/Contents")] = content_stream.indirect_reference

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_pdf_with_text(label: str = "Hello native text") -> bytes:
    """Construct a one-page PDF whose native text is non-empty."""

    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    writer._add_object(font)
    font_ref = font.indirect_reference

    resources = DictionaryObject()
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_ref
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources

    encoded = label.encode("latin-1")
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 50 50 Td (" + encoded + b") Tj ET")
    writer._add_object(content)
    page[NameObject("/Contents")] = content.indirect_reference

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class TestPdfOcrOptions:
    def test_normalized_clamps_out_of_range_values(self):
        normalized = PdfOcrOptions(
            enabled=True,
            language="",
            dpi=10_000,
            min_native_chars=0,
            max_pages=-5,
            timeout_seconds=0.1,
        ).normalized()
        assert normalized.dpi == 600
        assert normalized.min_native_chars == 1
        assert normalized.max_pages == 1
        assert normalized.timeout_seconds == 1.0
        assert normalized.language == "chi_sim+eng"


class TestPdfParserOcr:
    def test_native_text_page_does_not_invoke_ocr(self):
        pdf_bytes = _make_pdf_with_text(
            "Hello native text on this PDF page that is long enough"
        )

        renderer_called = {"value": False}
        tesseract_called = {"value": False}

        def renderer_factory(_pdf, _options):
            def _render(_page):
                renderer_called["value"] = True
                raise AssertionError("render should not be called")
            return _render

        def tesseract_runner(_png, _options):
            tesseract_called["value"] = True
            raise AssertionError("tesseract should not be called")

        parser = PdfParser(
            options=PdfOcrOptions(enabled=True),
            renderer_factory=renderer_factory,
            tesseract_runner=tesseract_runner,
            tesseract_lookup=lambda: "/usr/bin/tesseract",
        )
        result = parser.parse(pdf_bytes)
        assert result.success is True
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == []
        assert extraction["ocr_completed_pages"] == []
        assert extraction["native_text_pages"] == [1]
        assert not renderer_called["value"]
        assert not tesseract_called["value"]

    def test_scanned_candidate_page_runs_ocr_via_injected_renderer(self):
        pdf_bytes = _make_pdf_with_image()

        def renderer_factory(_pdf, _options):
            def _render(_page):
                from apps.api.parsers.pdf import _PageRender

                return _PageRender(
                    width_px=200,
                    height_px=200,
                    page_width_pt=200.0,
                    page_height_pt=200.0,
                    png=b"\x89PNG\r\n\x1a\nfake-bytes",
                )

            return _render

        tsv = (
            b"level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\t"
            b"left\ttop\twidth\theight\tconf\ttext\n"
            b"1\t1\t0\t0\t0\t0\t0\t0\t200\t200\t-1\t\n"
            b"5\t1\t1\t1\t1\t1\t10\t20\t40\t20\t92\thello\n"
            b"5\t1\t1\t1\t1\t2\t60\t20\t50\t20\t90\tworld\n"
        )

        def tesseract_runner(_png, _options):
            from subprocess import CompletedProcess

            return CompletedProcess(args=[], returncode=0, stdout=tsv, stderr=b"")

        parser = PdfParser(
            options=PdfOcrOptions(enabled=True, language="eng", min_native_chars=24),
            renderer_factory=renderer_factory,
            tesseract_runner=tesseract_runner,
            tesseract_lookup=lambda: "/usr/bin/tesseract",
        )
        result = parser.parse(pdf_bytes)
        assert result.success is True
        assert result.structured_content.full_text() == "hello world"
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == [1]
        assert extraction["ocr_completed_pages"] == [1]
        assert extraction["ocr_failed_pages"] == []
        assert extraction["ocr_status"] == "completed"
        ocr_block = result.structured_content.blocks[0]
        assert ocr_block.extra["source"] == "ocr"
        assert ocr_block.extra["ocr_engine"] == "tesseract"
        assert ocr_block.extra["bbox"] == [10.0, 160.0, 110.0, 180.0]
        assert 0 < ocr_block.extra["confidence"] <= 1

    def test_single_page_ocr_failure_still_succeeds(self):
        from subprocess import CompletedProcess

        from pypdf import PdfReader, PdfWriter

        from apps.api.parsers.pdf import _PageRender

        # Two image-only pages so we can verify one bad page does not abort.
        writer = PdfWriter()
        writer.add_page(PdfReader(BytesIO(_make_pdf_with_image())).pages[0])
        writer.add_page(PdfReader(BytesIO(_make_pdf_with_image())).pages[0])
        source = BytesIO()
        writer.write(source)

        def renderer_factory(_pdf, _options):
            def _render(_page):
                return _PageRender(
                    width_px=200,
                    height_px=200,
                    page_width_pt=200.0,
                    page_height_pt=200.0,
                    png=b"\x89PNG\r\n\x1a\nfake",
                )

            return _render

        calls = {"n": 0}

        def tesseract_runner(_png, _options):
            calls["n"] += 1
            if calls["n"] == 1:
                return CompletedProcess(args=[], returncode=1, stdout=b"", stderr=b"")
            tsv = (
                b"level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\t"
                b"left\ttop\twidth\theight\tconf\ttext\n"
                b"5\t1\t1\t1\t1\t1\t5\t5\t40\t20\t88\trecovered\n"
            )
            return CompletedProcess(args=[], returncode=0, stdout=tsv, stderr=b"")

        parser = PdfParser(
            options=PdfOcrOptions(enabled=True, language="eng", min_native_chars=24),
            renderer_factory=renderer_factory,
            tesseract_runner=tesseract_runner,
            tesseract_lookup=lambda: "/usr/bin/tesseract",
        )
        result = parser.parse(source.getvalue())
        assert result.success is True
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == [1, 2]
        assert extraction["ocr_completed_pages"] == [2]
        assert extraction["ocr_failed_pages"] == [1]
        assert result.structured_content.full_text() == "recovered"

    def test_disabled_options_record_skipped_pages(self):
        pdf_bytes = _make_pdf_with_image()

        parser = PdfParser(options=PdfOcrOptions(enabled=False))
        result = parser.parse(pdf_bytes)
        assert result.success is True
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == [1]
        assert extraction["ocr_skipped_pages"] == [1]
        assert extraction["ocr_completed_pages"] == []
        assert extraction["ocr_status"] == "skipped"
        reasons = extraction.get("ocr_skipped_reasons", {})
        assert "disabled" in reasons

    def test_missing_dependencies_record_skipped_pages(self, monkeypatch):
        pdf_bytes = _make_pdf_with_image()

        monkeypatch.setattr("apps.api.parsers.pdf.shutil.which", lambda _name: None)

        parser = PdfParser(
            options=PdfOcrOptions(enabled=True),
            tesseract_lookup=lambda: None,
        )
        result = parser.parse(pdf_bytes)
        assert result.success is True
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == [1]
        assert extraction["ocr_skipped_pages"] == [1]
        assert extraction["ocr_completed_pages"] == []
        assert extraction["ocr_status"] == "skipped"
        reasons = extraction.get("ocr_skipped_reasons", {})
        assert "dependencies_unavailable" in reasons

    def test_missing_dependencies_when_pypdfium2_missing(self, monkeypatch):
        import builtins

        original_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "pypdfium2" or name.startswith("pypdfium2."):
                raise ImportError("pypdfium2 not available in test env")
            return original_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)
        monkeypatch.setattr("apps.api.parsers.pdf.shutil.which", lambda _name: "/usr/bin/tesseract")

        pdf_bytes = _make_pdf_with_image()
        parser = PdfParser(options=PdfOcrOptions(enabled=True))
        result = parser.parse(pdf_bytes)
        assert result.success is True
        extraction = result.structured_content.metadata["pdf_extraction"]
        assert extraction["ocr_candidate_pages"] == [1]
        assert extraction["ocr_skipped_pages"] == [1]
        assert extraction["ocr_completed_pages"] == []
        reasons = extraction.get("ocr_skipped_reasons", {})
        assert "dependencies_unavailable" in reasons

    def test_ocr_disabled_does_not_require_pypdfium2(self, monkeypatch):
        import builtins

        original_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "pypdfium2" or name.startswith("pypdfium2."):
                raise ImportError("pypdfium2 not available in test env")
            return original_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)
        monkeypatch.setattr("apps.api.parsers.pdf.shutil.which", lambda _name: None)

        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        source = BytesIO()
        writer.write(source)

        parser = PdfParser(options=PdfOcrOptions(enabled=False))
        result = parser.parse(source.getvalue())
        assert result.success is True
        assert result.structured_content.metadata["page_count"] == 1


class TestGetParserForContentOcr:
    def test_pdf_parser_receives_options_via_get_parser(self):
        from apps.api.parsers.base import get_parser_for_content
        from apps.api.parsers.pdf import PdfOcrOptions

        options = PdfOcrOptions(enabled=True, language="eng", max_pages=7)
        parser = get_parser_for_content(
            "application/pdf", "scan.pdf", pdf_ocr_options=options
        )
        assert isinstance(parser, PdfParser)
        assert parser._options.language == "eng"
        assert parser._options.max_pages == 7

    def test_get_parser_for_content_without_options_remains_compatible(self):
        from apps.api.parsers.base import get_parser_for_content

        parser = get_parser_for_content("application/pdf", "doc.pdf")
        assert isinstance(parser, PdfParser)
        assert parser._options.enabled is False
        assert parser._options.language == "chi_sim+eng"

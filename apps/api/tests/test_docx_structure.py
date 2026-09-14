"""Tests for ``DocxParser`` body-order parsing (CZ-Q06).

These tests cover the bug fix where ``DocxParser.parse`` first emitted
every paragraph and only afterwards every table. The parser now walks
the document body's direct ``w:p``/``w:tbl`` children in document
order, so interleaved tables land under the right heading, consecutive
tables stay together, and parent section grouping downstream
(``build_chunk_specs``) keeps the correct attribution.

Limitations intentionally documented here:

* Only direct body children are visited. Nested tables and revision
  wrappers (``w:sdt``/``w:ins``/``w:del``/``w:customXml``) are not
  expanded; their inner blocks are not turned into separate records.
* Full table metadata (sheet name, cell coordinates, ...) is deferred
  to a later CZ-N step. ``metadata.docx_extraction`` only records the
  parser version, index semantics and a small capability surface.

The tests build synthetic DOCX in memory with ``python-docx`` so they
do not depend on a sample file in the repository.
"""

from __future__ import annotations

import io
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from apps.api.parsers import DocxParser
from apps.api.services.chunker import build_chunk_specs


DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)


def _save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_table(document: Document, *rows: tuple[str, ...]) -> None:
    if not rows:
        document.add_table(rows=1, cols=1)
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value


class TestDocxBodyOrder:
    def test_interleaved_paragraphs_tables_with_multiple_headings(self):
        document = Document()
        document.add_heading("总则", level=1)
        document.add_paragraph("总则第一段。")
        _add_table(document, ("项目", "数值"), ("收入", "100"))
        document.add_paragraph("总则第二段。")
        document.add_heading("细则", level=2)
        document.add_paragraph("细则第一段。")
        _add_table(document, ("注",),)
        document.add_paragraph("细则第二段。")
        raw = _save(document)

        result = DocxParser().parse(raw)

        assert result.success is True
        types = [block.type for block in result.structured_content.blocks]
        assert types == [
            "heading",
            "paragraph",
            "table",
            "paragraph",
            "heading",
            "paragraph",
            "table",
            "paragraph",
        ]
        indices = [block.paragraph_index for block in result.structured_content.blocks]
        assert indices == [0, 1, 2, 3, 4, 5, 6, 7]
        assert [block.extra["docx_body_index"] for block in result.structured_content.blocks] == indices

        assert result.structured_content.blocks[0].text == "总则"
        assert result.structured_content.blocks[0].heading_path == ["总则"]
        assert result.structured_content.blocks[2].text == "项目\t数值\n收入\t100"
        assert result.structured_content.blocks[2].heading_path == ["总则"]
        assert result.structured_content.blocks[6].text == "注"
        assert result.structured_content.blocks[6].heading_path == ["总则", "细则"]

    def test_table_between_headings_inherits_correct_path(self):
        document = Document()
        document.add_heading("章", level=1)
        document.add_paragraph("章正文。")
        document.add_heading("节", level=2)
        document.add_paragraph("节正文。")
        _add_table(document, ("字段", "值"))
        document.add_heading("小节", level=3)
        document.add_paragraph("小节正文。")
        _add_table(document, ("k", "v"))
        raw = _save(document)

        result = DocxParser().parse(raw)

        assert [block.heading_path for block in result.structured_content.blocks] == [
            ["章"],
            ["章"],
            ["章", "节"],
            ["章", "节"],
            ["章", "节"],
            ["章", "节", "小节"],
            ["章", "节", "小节"],
            ["章", "节", "小节"],
        ]
        # Tables sit at the right zero-based body index (章=0, 章正文=1,
        # 节=2, 节正文=3, 第一个表=4, 小节=5, 小节正文=6, 第二个表=7).
        tables = [b for b in result.structured_content.blocks if b.type == "table"]
        assert [b.paragraph_index for b in tables] == [4, 7]
        assert tables[0].text == "字段\t值"
        assert tables[1].text == "k\tv"

    def test_consecutive_tables_preserve_order_and_indices(self):
        document = Document()
        document.add_heading("对比", level=1)
        document.add_paragraph("两表对比。")
        _add_table(document, ("甲",))
        _add_table(document, ("乙",))
        _add_table(document, ("丙",))
        raw = _save(document)

        result = DocxParser().parse(raw)

        types = [block.type for block in result.structured_content.blocks]
        assert types == ["heading", "paragraph", "table", "table", "table"]
        assert [block.paragraph_index for block in result.structured_content.blocks] == [0, 1, 2, 3, 4]
        assert [block.text for block in result.structured_content.blocks if block.type == "table"] == ["甲", "乙", "丙"]
        for block in result.structured_content.blocks:
            assert block.extra["docx_body_index"] == block.paragraph_index
        assert result.structured_content.blocks[2].heading_path == ["对比"]
        assert result.structured_content.blocks[4].heading_path == ["对比"]

    def test_blank_paragraphs_count_toward_index_but_emit_no_block(self):
        document = Document()
        document.add_heading("章", level=1)
        document.add_paragraph("前文。")
        document.add_paragraph("")
        document.add_paragraph("   ")
        document.add_paragraph("后文。")
        _add_table(document, ("数据",))
        document.add_paragraph("")
        document.add_paragraph("收尾。")
        raw = _save(document)

        result = DocxParser().parse(raw)

        types = [block.type for block in result.structured_content.blocks]
        # The two blank paragraphs before and the blank after the table are
        # silently skipped; the surrounding paragraphs and the table
        # keep the correct body index.
        assert types == [
            "heading",
            "paragraph",
            "paragraph",
            "table",
            "paragraph",
        ]
        assert [block.paragraph_index for block in result.structured_content.blocks] == [
            0,
            1,
            4,
            5,
            7,
        ]
        assert [block.extra["docx_body_index"] for block in result.structured_content.blocks] == [
            0,
            1,
            4,
            5,
            7,
        ]
        assert result.structured_content.metadata["docx_extraction"][
            "body_indexed_children"
        ] == 8

    def test_document_with_only_tables_returns_table_blocks_in_order(self):
        document = Document()
        _add_table(document, ("第一", "列"))
        _add_table(document, ("第二",))
        raw = _save(document)

        result = DocxParser().parse(raw)

        assert [block.type for block in result.structured_content.blocks] == [
            "table",
            "table",
        ]
        assert [block.paragraph_index for block in result.structured_content.blocks] == [0, 1]
        assert result.structured_content.blocks[0].text == "第一\t列"
        assert result.structured_content.blocks[1].text == "第二"
        # No headings were ever seen, so heading_path stays empty.
        for block in result.structured_content.blocks:
            assert block.heading_path == []

    def test_empty_document_emits_no_blocks_but_records_index_count(self):
        # The default template contains sectPr, not an implicit paragraph.
        raw = _save(Document())

        result = DocxParser().parse(raw)

        assert result.success is True
        assert result.structured_content.blocks == []
        assert result.structured_content.metadata["docx_extraction"][
            "body_indexed_children"
        ] == 0

    def test_repeated_parse_is_deterministic(self):
        document = Document()
        document.add_heading("H1", level=1)
        document.add_paragraph("P1")
        _add_table(document, ("cell",))
        document.add_heading("H2", level=2)
        document.add_paragraph("P2")
        raw = _save(document)

        first = DocxParser().parse(raw)
        second = DocxParser().parse(raw)

        assert first.success is True and second.success is True
        first_dicts = [block.to_dict() for block in first.structured_content.blocks]
        second_dicts = [block.to_dict() for block in second.structured_content.blocks]
        assert first_dicts == second_dicts
        assert first.structured_content.full_text() == second.structured_content.full_text()
        assert first.structured_content.metadata == second.structured_content.metadata

    def test_parse_does_not_mutate_original_file_bytes(self):
        document = Document()
        document.add_heading("标题", level=1)
        document.add_paragraph("正文。")
        _add_table(document, ("列",))
        raw = _save(document)
        snapshot = memoryview(raw).tobytes()

        result = DocxParser().parse(raw)

        assert result.success is True
        assert raw == snapshot
        assert raw != b""

    def test_vendor_specific_style_type_still_extracts_heading(self):
        # Re-uses the malformed styles.xml trick from test_parsers.py to
        # make sure CZ-Q06 keeps the existing safe heading fallback.
        document = Document()
        document.add_heading("非标准样式标题", level=1)
        document.add_paragraph("正文仍应成功提取。")
        _add_table(document, ("值",))
        source = io.BytesIO()
        document.save(source)

        malformed = io.BytesIO()
        with (
            ZipFile(io.BytesIO(source.getvalue())) as input_archive,
            ZipFile(malformed, "w", ZIP_DEFLATED) as output_archive,
        ):
            for item in input_archive.infolist():
                data = input_archive.read(item.filename)
                if item.filename == "word/styles.xml":
                    data = data.replace(
                        b'w:type="paragraph" w:styleId="Heading1"',
                        b'w:type="titleLevel1" w:styleId="Heading1"',
                    )
                output_archive.writestr(item, data)

        result = DocxParser().parse(malformed.getvalue())

        assert result.success is True
        types = [block.type for block in result.structured_content.blocks]
        assert types == ["heading", "paragraph", "table"]
        assert result.structured_content.blocks[0].text == "非标准样式标题"
        assert result.structured_content.blocks[0].heading_path == ["非标准样式标题"]
        assert result.structured_content.blocks[2].heading_path == ["非标准样式标题"]
        assert result.structured_content.blocks[1].text == "正文仍应成功提取。"
        assert result.structured_content.blocks[2].text == "值"


class TestDocxExtractionMetadata:
    def test_metadata_exposes_parser_version_and_index_semantics(self):
        document = Document()
        document.add_heading("章", level=1)
        document.add_paragraph("段落。")
        _add_table(document, ("格",))
        raw = _save(document)

        result = DocxParser().parse(raw)

        extraction = result.structured_content.metadata.get("docx_extraction")
        assert isinstance(extraction, dict)
        assert extraction["parser_version"] == "1"
        assert isinstance(extraction["index_semantics"], str)
        assert "docx_body_index" in extraction["index_semantics"]
        # Capability surface: full table metadata is intentionally
        # out of scope for CZ-Q06, but the surface itself is part of
        # the documented limitation.
        assert extraction["nested_tables_supported"] is False
        assert extraction["revision_tracking_supported"] is False
        assert extraction["body_indexed_children"] == 3

    def test_metadata_capability_flags_advertise_limitations(self):
        raw = _save(Document())
        result = DocxParser().parse(raw)

        extraction = result.structured_content.metadata["docx_extraction"]
        assert "嵌套表格" in extraction["index_semantics"] or "nested" in extraction["index_semantics"].lower()


class TestDocxParentSectionAssociation:
    def test_build_chunk_specs_groups_tables_under_correct_section(self):
        document = Document()
        document.add_heading("总则", level=1)
        document.add_paragraph("总则正文。")
        _add_table(document, ("总则表行",))
        document.add_heading("细则", level=1)
        document.add_paragraph("细则正文。")
        _add_table(document, ("细则表行",))
        raw = _save(document)

        result = DocxParser().parse(raw)
        specs = build_chunk_specs(result.structured_content)
        parents = [spec for spec in specs if spec.parent_external_id is None]

        assert [parent.heading_path for parent in parents] == [["总则"], ["细则"]]
        assert "总则表行" in parents[0].content
        assert "细则表行" in parents[1].content
        assert parents[0].order_index < parents[1].order_index
        # Each parent also records the first body index it owns.
        assert parents[0].paragraph_index == 0
        assert parents[1].paragraph_index == 3

    def test_build_chunk_specs_keeps_table_after_subheading_in_parent_section(self):
        # A level-2 heading does not open a new section (chunker only
        # splits at level <= 1). The table that follows the subheading
        # should therefore live under the same parent as the
        # subheading text.
        document = Document()
        document.add_heading("章", level=1)
        document.add_paragraph("章正文。")
        document.add_heading("节", level=2)
        document.add_paragraph("节正文。")
        _add_table(document, ("节内表",))
        raw = _save(document)

        result = DocxParser().parse(raw)
        specs = build_chunk_specs(result.structured_content)
        parents = [spec for spec in specs if spec.parent_external_id is None]

        assert len(parents) == 1
        assert parents[0].content.startswith("章\n\n章正文。")
        joined = "\n".join(parent.content for parent in parents)
        assert "节" in joined
        assert "节内表" in joined
        # The single section spans the whole body, so the section's
        # start_paragraph_index matches the first heading.
        assert parents[0].paragraph_index == 0

    def test_table_between_two_sections_is_attributed_to_first_section(self):
        document = Document()
        document.add_heading("第一部分", level=1)
        document.add_paragraph("第一部分正文。")
        _add_table(document, ("第一表",))
        document.add_heading("第二部分", level=1)
        document.add_paragraph("第二部分正文。")
        raw = _save(document)

        result = DocxParser().parse(raw)
        tables = [block for block in result.structured_content.blocks if block.type == "table"]
        assert len(tables) == 1
        assert tables[0].text == "第一表"
        assert tables[0].heading_path == ["第一部分"]

        specs = build_chunk_specs(result.structured_content)
        parents = [spec for spec in specs if spec.parent_external_id is None]
        assert [parent.heading_path for parent in parents] == [["第一部分"], ["第二部分"]]
        assert "第一表" in parents[0].content
        assert "第二部分正文" in parents[1].content
        assert "第一表" not in parents[1].content


class TestDocxParserApi:
    def test_body_iteration_does_not_require_new_docx_helper(self, monkeypatch):
        from docx.oxml.document import CT_Body

        document = Document()
        document.add_paragraph("前文")
        _add_table(document, ("表格",))
        document.add_paragraph("后文")
        raw = _save(document)
        # Requirements permit python-docx 0.8, before inner_content_elements.
        monkeypatch.delattr(CT_Body, "inner_content_elements", raising=False)
        result = DocxParser().parse(raw)
        assert result.success
        assert result.structured_content.full_text() == "前文\n表格\n后文"

    def test_table_stays_between_explanation_and_note_in_chunks(self):
        document = Document()
        document.add_heading("监测结果", level=1)
        document.add_paragraph("前置说明")
        _add_table(document, ("项目", "单位"), ("总量", "千克"))
        document.add_paragraph("表下注释")
        structured = DocxParser().parse(_save(document)).structured_content
        specs = build_chunk_specs(structured)
        parent = next(spec for spec in specs if spec.role == "parent")
        assert parent.content.index("前置说明") < parent.content.index("总量")
        assert parent.content.index("总量") < parent.content.index("表下注释")
        children = [spec for spec in specs if spec.role == "child"]
        assert all(spec.parent_external_id == parent.external_id for spec in children)
        assert any("总量" in spec.content for spec in children)
        # Stored JSON and the in-memory structure follow the same chunk path.
        assert specs == build_chunk_specs(structured.to_dict())

    def test_can_parse_matches_mime_and_extension(self):
        parser = DocxParser()
        assert parser.can_parse(DOCX_MIME, "anything.docx") is True
        assert parser.can_parse(None, "report.docx") is True
        assert parser.can_parse("application/octet-stream", "report.docx") is True
        assert parser.can_parse("application/pdf", "report.pdf") is False
        assert parser.can_parse(None, None) is False

    def test_import_error_returns_clean_failure(self, monkeypatch):
        import builtins

        original_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "docx" or name.startswith("docx."):
                raise ImportError("docx disabled in test")
            return original_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)

        result = DocxParser().parse(b"not a real docx")

        assert result.success is False
        assert "python-docx" in (result.error_message or "")
        assert result.error_details == {"import_error": "python-docx not found"}
